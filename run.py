import os
import sys
import re
import time
import traceback
from datetime import datetime
import pytz
from playwright.sync_api import sync_playwright
from pydantic import BaseModel, Field
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont

# ---------------------------------------------------------
# DATA SCHEMAS
# ---------------------------------------------------------
class CauseListRow(BaseModel):
    sl_no: str = Field(description="Serial number")
    case_number: str = Field(description="Case number like WP 6554/2026")
    case_name: str = Field(description="Party names and representation details")
    ch: str = Field(description="Court Hall number")
    list_num: str = Field(description="List number")
    list_sl_no: str = Field(description="Item number in the list")
    status: str = Field(description="Case status like ORDERS")
    judges: str = Field(description="Hon'ble Judge name(s)")

class CauseListDocument(BaseModel):
    date: str = Field(description="Date of the cause list")
    rows: list[CauseListRow]

# ---------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------
def get_target_date_ist():
    manual = os.environ.get("TEST_DATE")
    if manual and manual.strip():
        return manual.strip()
    ist = pytz.timezone("Asia/Kolkata")
    return datetime.now(ist).strftime("%d/%m/%Y")

def clean_party_string(s: str) -> str:
    """Strips out leading PET/RES prefixes, procedural notes, and advocate names."""
    if not s:
        return ""
    
    # 1. Strip leading party tags (PET:, RES:, etc.)
    s = re.sub(r'^\s*(?:PET|RES|PETITIONER|RESPONDENT|APPELLANT|COMPLAINANT|APPLICANT|RESPODNENT)\s*[:.]?\s*', '', s, flags=re.I)
    
    # 2. Strip procedural bracket notes like (GM, RES), (SC, ), (DATE), (CH MOVED)
    s = re.sub(r'\([^\)]*\)', ' ', s)
    
    # 3. Strip trailing procedural notes
    s = re.split(r'\b(?:A/W|REG\s*:|V/O\s+DTD|V/O/D|MEMO\s+FOR|NOTE\s*:|OFFICE\s+OBJ)\b', s, flags=re.I)[0]
    
    # 4. Cut off at explicit advocate delimiters
    adv_delim = re.split(r'\b(?:ADV|ADVOCATE|ADVOCATES|COUNSEL|GOVT\s+ADVOCATE|HCGP|AGA|SPP|C/R\d*|PARTY\s+IN\s+PERSON)\s*[:.]?\s*', s, flags=re.I)
    s = adv_delim[0]
    
    # 5. Remove lawyer tokens and trailing tags
    tokens = [
        r'\bMAHESH\s+CHOWDHA?R[YI]\b',
        r'\bMAHESH\s+C\b',
        r'\bM\s+CHOWDHA?R[YI]\b',
        r'\bAGA\b',
        r'\bHCGP\b',
        r'\bSPP\b',
        r'\bGA\s+SD\b',
        r'\bSD\b',
        r'\bFOR\s+(?:RES|PET|R\d+|C/R)\b',
        r'\bVK\s+NOT\s+FILED\b'
    ]
    for pat in tokens:
        s = re.sub(pat, '', s, flags=re.I)
        
    s = re.sub(r'\s+', ' ', s).strip(" ,;:-./")
    return s.upper()

def clean_case_details(raw_name: str):
    """
    Splits case text into Petitioner and Respondent,
    and identifies which side Advocate Mahesh Chowdhary represents.
    """
    text = re.sub(r'[\r\n]+', ' ', raw_name)
    text = re.sub(r'\s+', ' ', text).strip()

    vs_match = re.search(r'\s+(?:V/?S\.?|VERSUS)\s+', text, re.I)
    res_match = re.search(r'\b(?:RES|RESPONDENT|RESPODNENT)\s*:\s*', text, re.I)

    if vs_match:
        pet_block = text[:vs_match.start()].strip()
        res_block = text[vs_match.end():].strip()
    elif res_match:
        pet_block = text[:res_match.start()].strip()
        res_block = text[res_match.start():].strip()
    else:
        pet_block = text
        res_block = ""

    # Detect which party Mahesh Chowdhary represents
    if re.search(r'\b(MAHESH|CHOWDHA?R[YI]|CHOUDHARY)\b', res_block, re.I):
        in_charge = "RES"
    else:
        in_charge = "PET"

    pet_clean = clean_party_string(pet_block)
    res_clean = clean_party_string(res_block)

    return pet_clean, res_clean, in_charge

# ---------------------------------------------------------
# SCRAPING STEP (PLAYWRIGHT)
# ---------------------------------------------------------
def fetch_court_pdf(pdf_path="causelist.pdf", text_path="causelist_text.txt"):
    date_str = get_target_date_ist()
    print(f"[*] Target Causelist Date: {date_str}")

    scraperapi_key = os.environ.get("SCRAPERAPI_KEY")
    launch_args = {
        "headless": True,
        "args": ["--no-sandbox", "--disable-setuid-sandbox", "--disable-dev-shm-usage", "--ignore-certificate-errors"]
    }
    if scraperapi_key and scraperapi_key.strip():
        launch_args["proxy"] = {
            "server": "http://proxy-server.scraperapi.com:8001",
            "username": "scraperapi.country_code=in",
            "password": scraperapi_key.strip()
        }

    with sync_playwright() as p:
        browser = p.chromium.launch(**launch_args)
        context = browser.new_context(
            ignore_https_errors=True,
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36",
            viewport={"width": 1366, "height": 900}
        )
        page = context.new_page()
        page.add_init_script("window.print = () => { console.log('print intercepted'); };")

        print("[*] Navigating to Karnataka High Court Portal...")
        page.goto("https://judiciary.karnataka.gov.in/causelistSearch.php", wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(3000)

        # 1. Bench Selection
        bench_select = page.locator("select[name='bench']:visible").first
        for opt in bench_select.locator("option").all():
            opt_text = opt.inner_text().strip()
            if "Bengaluru" in opt_text or "Bangalore" in opt_text or "Principal" in opt_text:
                bench_select.select_option(label=opt_text)
                break
        page.wait_for_timeout(1500)

        # 2. Search By: Advocate
        searchby_select = page.locator("select[name='searchby']:visible").first
        for opt in searchby_select.locator("option").all():
            if "Advocate" in opt.inner_text().strip():
                searchby_select.select_option(label=opt.inner_text().strip())
                break
        page.wait_for_timeout(1500)

        # 3. Enter Dates
        page.locator("#fromDt:visible").first.fill(date_str)
        page.locator("#toDt:visible").first.fill(date_str)
        page.wait_for_timeout(1000)

        # 4. Enter Advocate Name
        adv_input = page.locator("input[placeholder='Enter Advocate Name']:visible").first
        if not adv_input.is_visible():
            adv_input = page.locator("#advName:visible").first
        adv_input.click()
        adv_input.fill("")
        adv_input.fill("Mahesh Chowdhary")
        page.wait_for_timeout(1000)

        # 5. Click GET DETAILS
        page.locator("#getData:visible").first.click()
        print("[*] Waiting for search results...")
        page.wait_for_timeout(7000)

        # 6. Click Print Button
        print_btn = None
        for btn in page.locator("input[type='button']:visible, button:visible").all():
            val = (btn.get_attribute("value") or "").lower()
            txt = (btn.inner_text() or "").lower()
            if btn.get_attribute("id") == "getData":
                continue
            if "print" in val or "print" in txt:
                print_btn = btn
                break

        target_page = page
        if print_btn:
            try:
                with context.expect_page(timeout=8000) as popup_info:
                    print_btn.click()
                target_page = popup_info.value
                target_page.wait_for_load_state("domcontentloaded")
                target_page.wait_for_timeout(3000)
            except Exception:
                target_page = page

        # Save both PDF and Raw Page Text (for reliable manual fallback)
        target_page.pdf(path=pdf_path, format="A4", print_background=True)
        print(f"[+] Downloaded court PDF to {pdf_path}")
        
        page_text = target_page.inner_text("body")
        with open(text_path, "w", encoding="utf-8") as f:
            f.write(page_text)
        print(f"[+] Saved page text backup to {text_path}")

        browser.close()

# ---------------------------------------------------------
# GEMINI PARSER
# ---------------------------------------------------------
def parse_pdf_with_gemini(pdf_path="causelist.pdf"):
    """Uses Gemini API to parse cases."""
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key or not api_key.strip():
        raise ValueError("GEMINI_API_KEY not found in environment.")

    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    with open(pdf_path, "rb") as f:
        pdf_bytes = f.read()

    prompt = (
        "Extract the complete cause list table from this PDF into the structured JSON schema. "
        "Strictly extract only actual case rows with case numbers and parties. "
        "In case_name, preserve the full text of petitioner, respondent, and advocates "
        "(e.g., 'PET: ... ADV: ... RES: ... ADV: ...') so advocate representation can be verified. "
        "Do NOT extract banner notices or general instructions. "
        "If no cases are found, return an empty rows array."
    )

    models = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-1.5-flash"]
    for model_name in models:
        try:
            print(f"[*] Submitting PDF to Gemini ({model_name})...")
            response = client.models.generate_content(
                model=model_name,
                contents=[types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf"), prompt],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=CauseListDocument,
                    temperature=0.0
                )
            )
            parsed = CauseListDocument.model_validate_json(response.text)
            print(f"[+] Successfully extracted {len(parsed.rows)} cases using {model_name}.")
            return parsed.rows
        except Exception as e:
            print(f"[!] {model_name} failed ({e}), trying fallback model...")
            time.sleep(1)

    raise RuntimeError("All Gemini models failed or returned errors.")

# ---------------------------------------------------------
# MANUAL FALLBACK PARSER
# ---------------------------------------------------------
def parse_cause_list_manually(pdf_path="causelist.pdf", text_path="causelist_text.txt") -> list[CauseListRow]:
    """
    Manually extracts cases from text or PDF without LLM.
    Tracks Court Hall, Cause List No, Judge(s), Status, and Case rows.
    """
    print("[*] Running manual parser...")
    full_text = ""
    if os.path.exists(text_path):
        with open(text_path, "r", encoding="utf-8") as f:
            full_text = f.read()

    # Fallback to local PDF extraction via pypdf if text file is empty
    if not full_text.strip() and os.path.exists(pdf_path):
        try:
            import pypdf
            reader = pypdf.PdfReader(pdf_path)
            extracted = [page.extract_text() for page in reader.pages if page.extract_text()]
            full_text = "\n".join(extracted)
        except Exception as e:
            print(f"[!] pypdf extraction note: {e}")

    if not full_text.strip():
        print("[!] No cause list text available for manual parsing.")
        return []

    lines = [l.strip() for l in full_text.splitlines() if l.strip()]

    current_ch = "-"
    current_list = "1"
    current_judges = "HON'BLE HIGH COURT"
    current_status = "ORDERS"

    status_keywords = [
        "PRELIMINARY HEARING - B GROUP",
        "PRELIMINARY HEARING (B GROUP)",
        "PRELIMINARY HEARING",
        "HEARING - INTERLOCUTORY APPLN",
        "HEARING - INTERLOCUTORY APPLICATION",
        "HEARING - INTERLOCUTORY",
        "FURTHER HEARING",
        "FINAL HEARING",
        "ADMISSION",
        "ORDERS",
        "FOR ORDERS",
        "FOR ADMISSION",
        "FRESH MATTERS",
        "DICTATING ORDERS",
        "FINAL DISPOSAL"
    ]

    case_start_regex = re.compile(
        r'^\*?\s*(\d+)\s+((?:WP|CRL\.P|CRL\.A|WA|MFA|CCC|CRP|RSA|RFA|HRRP|CP|EP|CEA|STA|ITA|WTA|RPFC|WP\(C\)|CONT\.P)\s*(?:NO\.?)?\s*\d+/\d{2,4})\b',
        re.I
    )

    rows = []
    accumulating_case = None

    def flush_case(acc):
        if not acc:
            return
        rows.append(CauseListRow(
            sl_no=str(len(rows) + 1),
            case_number=acc["case_number"].strip(),
            case_name=acc["case_name"].strip(),
            ch=acc["ch"],
            list_num=acc["list_num"],
            list_sl_no=acc["list_sl_no"],
            status=acc["status"],
            judges=acc["judges"]
        ))

    for line in lines:
        # 1. Court Hall
        ch_match = re.search(r'COURT\s+HALL\s*(?:NO\.?|NUMBER)?\s*[:\-]?\s*([0-9A-Z]+)', line, re.I)
        if ch_match:
            current_ch = ch_match.group(1).strip()
            continue

        # 2. List Number
        list_match = re.search(r'(?:Cause\s+List|List)\s*(?:NO\.?|NUMBER)?\s*[:\-]?\s*(\d+)', line, re.I)
        if list_match:
            current_list = list_match.group(1).strip()
            continue

        # 3. Judges
        judge_match = re.search(r'(?:BEFORE\s+)?(HON(?:[\'`\u2019])?BLE\s+(?:MR\.|MRS\.|MS\.)?\s*JUSTICE\s+[A-Z\.\s]+)', line, re.I)
        if judge_match:
            current_judges = judge_match.group(1).replace("`", "'").replace("’", "'").strip().upper()
            continue

        # 4. Status
        matched_status = False
        for sk in status_keywords:
            if re.search(rf'^{re.escape(sk)}\b', line, re.I):
                current_status = sk.upper()
                matched_status = True
                break
        if matched_status:
            continue

        # 5. Case row detection
        case_match = case_start_regex.match(line)
        if case_match:
            flush_case(accumulating_case)
            list_item_no = case_match.group(1).strip()
            case_no = case_match.group(2).strip().upper()
            remaining_text = line[case_match.end():].strip()

            accumulating_case = {
                "list_sl_no": list_item_no,
                "case_number": case_no,
                "case_name": remaining_text,
                "ch": current_ch,
                "list_num": current_list,
                "status": current_status,
                "judges": current_judges
            }
            continue

        # If we are currently collecting lines for a case
        if accumulating_case:
            if re.search(r'^(?:IN\s+THE\s+HIGH\s+COURT|COURT\s+HALL|PHYSICAL\s+HEARING|BEFORE|Cause\s+List)', line, re.I):
                flush_case(accumulating_case)
                accumulating_case = None
            else:
                accumulating_case["case_name"] += " " + line

    flush_case(accumulating_case)
    return rows

# ---------------------------------------------------------
# UNIFIED EXTRACTOR (GEMINI -> MANUAL FALLBACK)
# ---------------------------------------------------------
def extract_cause_list_rows(pdf_path="causelist.pdf", text_path="causelist_text.txt") -> list[CauseListRow]:
    rows = []
    gemini_key = os.environ.get("GEMINI_API_KEY")

    if gemini_key and gemini_key.strip():
        try:
            print("[*] Attempting extraction with Gemini API...")
            rows = parse_pdf_with_gemini(pdf_path=pdf_path)
            if rows:
                return rows
            print("[!] Gemini returned 0 rows. Falling back to manual parser...")
        except Exception as e:
            print(f"[!] Gemini unavailable ({e}). Automatically switching to manual parser...")
    else:
        print("[*] GEMINI_API_KEY not configured. Running manual parser directly...")

    rows = parse_cause_list_manually(pdf_path=pdf_path, text_path=text_path)
    return rows

# ---------------------------------------------------------
# EXCEL GENERATION (EXACT MATCH TO ATTACHED IMAGE)
# ---------------------------------------------------------
def write_to_excel(rows: list[CauseListRow], excel_path="cause_list.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cause List"
    ws.views.sheetView[0].showGridLines = True

    # 1. Header setup
    headers = ["SL NO", "CASE NUMBER", "CASE NAME", "CH", "LIST", "SL NO", "STATUS", "JUDGES"]
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    header_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10, bold=False)
    bold_cell_font = Font(name="Calibri", size=10, bold=True)

    inline_bold = InlineFont(rFont="Calibri", sz=10, b=True)
    inline_regular = InlineFont(rFont="Calibri", sz=10, b=False)

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000")
    )

    for col_idx in range(1, 9):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # 2. Add Data Rows
    for idx, row in enumerate(rows, start=1):
        pet_clean, res_clean, in_charge = clean_case_details(row.case_name)
        plain_text = f"{pet_clean}\nvs\n{res_clean}"

        ws.append([
            idx,
            row.case_number.strip().upper(),
            plain_text,
            row.ch.strip(),
            row.list_num.strip(),
            row.list_sl_no.strip(),
            row.status.strip().upper(),
            row.judges.strip().upper()
        ])

        curr_row = idx + 1
        ws.row_dimensions[curr_row].height = 48

        # 3. Rich Text Bolding for Mahesh Chowdhary's side
        case_cell = ws.cell(row=curr_row, column=3)
        try:
            if in_charge == "PET":
                case_cell.value = CellRichText(
                    TextBlock(inline_bold, pet_clean),
                    TextBlock(inline_regular, f"\nvs\n{res_clean}")
                )
            else:
                case_cell.value = CellRichText(
                    TextBlock(inline_regular, f"{pet_clean}\nvs\n"),
                    TextBlock(inline_bold, res_clean)
                )
        except Exception:
            case_cell.value = plain_text

    # 4. Cell alignments, borders, and bolding
    for r in range(2, ws.max_row + 1):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border

            if c in [1, 4, 5, 6]:
                # Center aligned normal text (SL NO, CH, LIST, SL NO)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = regular_font
            elif c == 2:
                # Center aligned BOLD Case Number
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_cell_font
            elif c == 3:
                # Left aligned Case Name with wrap text
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                # Left aligned Status and Judges with wrap text
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.font = regular_font

    col_widths = {1: 8, 2: 24, 3: 45, 4: 8, 5: 8, 6: 10, 7: 28, 8: 38}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    wb.save(excel_path)
    print(f"[+] Excel generated successfully at {excel_path}")

# ---------------------------------------------------------
# PDF REPORT GENERATION
# ---------------------------------------------------------
def generate_pdf(rows: list[CauseListRow], date_str: str, pdf_path="cause_list.pdf"):
    """Renders landscape black-and-white table PDF."""
    rows_html = ""
    for idx, row in enumerate(rows, start=1):
        pet_clean, res_clean, in_charge = clean_case_details(row.case_name)
        if in_charge == "PET":
            case_cell = f"<strong>{pet_clean}</strong><br>vs<br>{res_clean}"
        else:
            case_cell = f"{pet_clean}<br>vs<br><strong>{res_clean}</strong>"

        rows_html += f"""
        <tr>
            <td class="text-center">{idx}</td>
            <td class="text-center font-bold">{row.case_number}</td>
            <td>{case_cell}</td>
            <td class="text-center">{row.ch}</td>
            <td class="text-center">{row.list_num}</td>
            <td class="text-center">{row.list_sl_no}</td>
            <td>{row.status}</td>
            <td>{row.judges}</td>
        </tr>
        """

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @page {{ size: A4 landscape; margin: 10mm; }}
            body {{ font-family: Calibri, Arial, sans-serif; color: #000; margin: 0; font-size: 11px; }}
            .header {{ display: flex; justify-content: space-between; border-bottom: 1.5px solid #000; padding-bottom: 4px; margin-bottom: 8px; }}
            .header h1 {{ margin: 0; font-size: 14px; text-transform: uppercase; }}
            table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
            th {{ border: 1px solid #000; padding: 6px; font-weight: bold; text-align: center; background: #fff; font-size: 10px; }}
            td {{ border: 1px solid #000; padding: 6px; vertical-align: middle; word-wrap: break-word; font-size: 10px; }}
            .text-center {{ text-align: center; }}
            .font-bold {{ font-weight: bold; }}
        </style>
    </head>
    <body>
        <div class="header">
            <div>
                <h1>High Court of Karnataka &mdash; Bengaluru Bench</h1>
                <div>Daily Cause List &bull; Date: {date_str}</div>
            </div>
            <div style="text-align: right;">
                <div style="font-weight: bold;">ADVOCATE: MAHESH CHOWDHARY</div>
                <div>Total Matters: {len(rows)}</div>
            </div>
        </div>
        <table>
            <colgroup>
                <col style="width: 5%;">
                <col style="width: 15%;">
                <col style="width: 34%;">
                <col style="width: 5%;">
                <col style="width: 5%;">
                <col style="width: 6%;">
                <col style="width: 14%;">
                <col style="width: 16%;">
            </colgroup>
            <thead>
                <tr>
                    <th>SL NO</th>
                    <th>CASE NUMBER</th>
                    <th>CASE NAME</th>
                    <th>CH</th>
                    <th>LIST</th>
                    <th>SL NO</th>
                    <th>STATUS</th>
                    <th>JUDGES</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
    </body>
    </html>
    """

    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox", "--disable-setuid-sandbox"])
        page = browser.new_page()
        page.set_content(html_content, wait_until="networkidle")
        page.pdf(
            path=pdf_path,
            format="A4",
            landscape=True,
            print_background=True,
            margin={"top": "10mm", "bottom": "10mm", "left": "10mm", "right": "10mm"}
        )
        browser.close()
    print(f"[+] PDF cause list generated at {pdf_path}")

# ---------------------------------------------------------
# MAIN ORCHESTRATOR
# ---------------------------------------------------------
if __name__ == "__main__":
    try:
        raw_pdf = "causelist.pdf"
        raw_text = "causelist_text.txt"
        
        # 1. Scrape Court PDF and Raw Page Text
        fetch_court_pdf(raw_pdf, raw_text)
        target_date = get_target_date_ist()
        
        # 2. Extract Data (Gemini with automatic manual fallback)
        rows = extract_cause_list_rows(raw_pdf, raw_text)
        print(f"[+] Total {len(rows)} matters extracted.")
        
        # 3. Export to Excel & PDF
        write_to_excel(rows, "cause_list.xlsx")
        generate_pdf(rows, target_date, "cause_list.pdf")
        
        print("\n[SUCCESS] Both cause_list.xlsx and cause_list.pdf are ready in the requested format!")
    except Exception as err:
        print(f"[FATAL] Process aborted: {err}", file=sys.stderr)
        traceback.print_exc()
        sys.exit(1)
